# apps/attendance/exports/attendance_exports.py
from __future__ import annotations

import io
from typing import Any, Iterable, List, cast

from django.http import HttpResponse
from django.template.loader import get_template
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from weasyprint import HTML

__all__ = [
    "export_attendance_pdf",
    "export_attendance_excel",
    "export_lateness_absence_excel",
]

CONTENT_TYPE_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _name_from_record(obj: Any) -> str:
    ev = getattr(obj, "evaluation", None)
    if ev and getattr(ev, "employee", None) is not None:
        emp = ev.employee
        name = getattr(emp, "name", None)
        if name:
            return str(name)
    return ""


def _date_str(d: Any) -> str:
    return d.strftime("%Y-%m-%d") if (d is not None and hasattr(d, "strftime")) else ""


def _status_str(obj: Any) -> str:
    return str(obj.status) if hasattr(obj, "status") and obj.status is not None else ""


def export_attendance_pdf(
    records: Iterable[Any],
    template_path: str = "attendance/attendance_pdf.html",
    filename: str = "attendance.pdf",
) -> HttpResponse:
    html = get_template(template_path).render({"records": list(records)})
    pdf_bytes = HTML(string=html).write_pdf() or b""
    resp = HttpResponse(pdf_bytes, content_type="application/pdf")
    resp["Content-Disposition"] = f'attachment; filename="{filename}"'
    return resp


def _wb_response(wb: Workbook, filename: str) -> HttpResponse:
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    resp = HttpResponse(buf.read(), content_type=CONTENT_TYPE_XLSX)
    resp["Content-Disposition"] = f'attachment; filename="{filename}"'
    return resp


def export_attendance_excel(records: Iterable[Any]) -> HttpResponse:
    wb = Workbook()
    ws = cast(Worksheet, wb.active)
    if ws is None:  # pragma: no cover - stub safety
        ws = wb.create_sheet(title="Attendance")
    ws.title = "Attendance"
    headers: List[str] = ["اسم الموظف", "التاريخ", "الحالة"]
    ws.append(headers)
    for r in records:
        ws.append(
            [_name_from_record(r), _date_str(getattr(r, "date", None)), _status_str(r)]
        )
    return _wb_response(wb, "attendance.xlsx")


def export_lateness_absence_excel(entries: Iterable[Any]) -> HttpResponse:
    wb = Workbook()
    ws = cast(Worksheet, wb.active)
    if ws is None:  # pragma: no cover - stub safety
        ws = wb.create_sheet(title="Lateness & Absence")
    ws.title = "Lateness & Absence"
    headers: List[str] = ["اسم الموظف", "التاريخ", "الحالة"]
    ws.append(headers)
    for e in entries:
        ws.append(
            [_name_from_record(e), _date_str(getattr(e, "date", None)), _status_str(e)]
        )
    return _wb_response(wb, "lateness_absence.xlsx")
